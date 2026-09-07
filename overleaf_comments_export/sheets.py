"""The export as rows, for a spreadsheet.

Markdown is for reading and JSON is for machines. A sheet is for working
through a review: sort by file, filter to one reviewer, tick things off.

The shape is the hard part rather than the file format. Comments have replies,
so a single flat table forces a choice between losing the conversation and
repeating every comment's details on every reply row, and the second makes
filtering lie to you. Three sheets keyed on the short id avoid both.

This module only builds rows. Python writes them with openpyxl and the browser
extension writes them with its own small xlsx writer, and the parity test
compares what comes out of here against its JavaScript twin, so the two cannot
drift into different spreadsheets.
"""

from __future__ import annotations

from typing import Any

COMMENT_COLUMNS = [
    "Comment", "File", "Line", "Section", "Figure or table", "Commented on",
    "Author", "Raised", "Status", "Replies", "Comment text",
]
REPLY_COLUMNS = ["Comment", "Reply", "Author", "Written", "Reply text"]
CHANGE_COLUMNS = [
    "Change", "File", "Line", "Section", "Kind", "Author", "When", "Text",
]


def _float_label(enclosing: dict | None) -> str:
    """"Figure 3", or the caption when the number cannot be trusted."""
    if not enclosing:
        return ""
    kind = str(enclosing.get("kind") or "").capitalize()
    number = enclosing.get("number")
    if number:
        return f"{kind} {number}"
    label = enclosing.get("label")
    caption = enclosing.get("caption")
    if caption:
        return f"{kind} “{caption}”" if not label else f"{kind} “{caption}” ({label})"
    return f"{kind} (unnumbered)"


def _person(user: dict | None) -> str:
    user = user or {}
    return str(user.get("name") or user.get("email") or user.get("id") or "")


def _messages(thread: dict | None) -> list[dict]:
    if not isinstance(thread, dict):
        return []
    return [m for m in (thread.get("messages") or []) if isinstance(m, dict)]


def build_rows(payload: dict[str, Any]) -> dict[str, list[list[Any]]]:
    """Every sheet, as a header row followed by data rows."""
    threads = payload.get("threads") or {}

    comments: list[list[Any]] = [list(COMMENT_COLUMNS)]
    replies: list[list[Any]] = [list(REPLY_COLUMNS)]

    for c in payload.get("comments") or []:
        thread = threads.get(c.get("thread_id"))
        msgs = _messages(thread)
        first = msgs[0] if msgs else {}
        resolved = bool((thread or {}).get("resolved"))
        comments.append([
            c.get("short_id") or "",
            c.get("pathname") or "",
            c.get("line") or "",
            c.get("nearest_heading") or "",
            _float_label(c.get("enclosing_float")),
            c.get("anchored_text") or "",
            _person(first.get("user")),
            first.get("timestamp") or "",
            "Resolved" if resolved else "Open",
            max(0, len(msgs) - 1),
            first.get("content") or "",
        ])
        for i, m in enumerate(msgs[1:], start=1):
            replies.append([
                c.get("short_id") or "",
                i,
                _person(m.get("user")),
                m.get("timestamp") or "",
                m.get("content") or "",
            ])

    changes: list[list[Any]] = [list(CHANGE_COLUMNS)]
    for ch in payload.get("tracked_changes") or []:
        changes.append([
            ch.get("short_id") or "",
            ch.get("pathname") or "",
            ch.get("line") or "",
            ch.get("nearest_heading") or "",
            str(ch.get("kind") or "").capitalize(),
            _person(ch.get("user")),
            ch.get("timestamp") or "",
            ch.get("content") or "",
        ])

    return {"Comments": comments, "Replies": replies, "Tracked changes": changes}


def write_xlsx(payload: dict[str, Any], path) -> None:
    """Write the sheets to an .xlsx file.

    openpyxl is an optional extra rather than a dependency, because a
    spreadsheet is not what most people come here for and it is a large
    install to impose on everyone.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        from openpyxl.utils import get_column_letter
    except ImportError as e:  # pragma: no cover - exercised by the CLI path
        raise RuntimeError(
            "Writing a spreadsheet needs openpyxl. Install it with:\n\n"
            '    pip install "overleaf-comments-export[xlsx]"'
        ) from e

    sheets = build_rows(payload)
    book = Workbook()
    book.remove(book.active)

    for name, rows in sheets.items():
        sheet = book.create_sheet(title=name)
        for row in rows:
            sheet.append(row)
        # A header you can still read after scrolling, and filters, because
        # the entire point of this format is sorting and filtering.
        sheet.freeze_panes = "A2"
        if len(rows) > 1:
            sheet.auto_filter.ref = (
                f"A1:{get_column_letter(len(rows[0]))}{len(rows)}")
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for i, header in enumerate(rows[0], start=1):
            longest = max((len(str(r[i - 1])) for r in rows[1:]), default=0)
            # Comment text runs to paragraphs, so it is capped and wrapped
            # rather than allowed to make one column wider than the screen.
            width = min(max(len(str(header)) + 2, min(longest + 2, 60)), 60)
            sheet.column_dimensions[get_column_letter(i)].width = width
        text_columns = {"Comment text", "Reply text", "Text", "Commented on"}
        for i, header in enumerate(rows[0], start=1):
            if header in text_columns:
                for row in sheet.iter_rows(min_row=2, min_col=i, max_col=i):
                    for cell in row:
                        cell.alignment = Alignment(wrap_text=True, vertical="top")

    book.save(path)
