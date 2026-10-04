"""The spreadsheet's shape.

A sheet is flat and the data is not, so the shape is the whole decision. One
table would force a choice between losing the replies and repeating every
comment's details on every reply row, and the second makes filtering lie.
"""

from __future__ import annotations

import pytest

from overleaf_comments_export.sheets import (COMMENT_COLUMNS, build_rows,
                                             _float_label)


def payload(**over):
    base = {
        "threads": {
            "t1": {"resolved": False, "messages": [
                {"id": "m1", "content": "Break this up.",
                 "timestamp": "2026-08-01T10:00:00+00:00",
                 "user": {"name": "Bakhtawar Khan"}},
                {"id": "m2", "content": "Agreed.",
                 "timestamp": "2026-08-02T10:00:00+00:00",
                 "user": {"name": "ans.ahmad"}},
            ]},
            "t2": {"resolved": True, "messages": [
                {"id": "m3", "content": "Done.",
                 "timestamp": "2026-08-01T11:00:00+00:00",
                 "user": {"email": "someone@example.com"}}]},
        },
        "comments": [
            {"short_id": "C001", "thread_id": "t1", "pathname": "main.tex",
             "line": 12, "nearest_heading": "Method", "anchored_text": "this bit",
             "enclosing_float": {"kind": "figure", "number": 3,
                                 "label": "fig:x", "caption": "A plot"}},
            {"short_id": "C002", "thread_id": "t2", "pathname": "intro.tex",
             "line": 3, "nearest_heading": None, "anchored_text": "that bit",
             "enclosing_float": None},
        ],
        "tracked_changes": [
            {"short_id": "T001", "pathname": "main.tex", "line": 9,
             "nearest_heading": "Method", "kind": "insertion",
             "content": "very ", "timestamp": "2026-08-03T10:00:00+00:00",
             "user": {"name": "ans.ahmad"}},
        ],
    }
    base.update(over)
    return base


def test_three_sheets_with_headers():
    rows = build_rows(payload())
    assert list(rows) == ["Comments", "Replies", "Tracked changes"]
    assert rows["Comments"][0] == COMMENT_COLUMNS


def test_one_row_per_comment_not_per_message():
    rows = build_rows(payload())
    assert len(rows["Comments"]) == 3          # header plus two comments
    assert [r[0] for r in rows["Comments"][1:]] == ["C001", "C002"]


def test_replies_are_their_own_sheet_keyed_back_to_the_comment():
    """So a reply can be read without repeating the whole comment beside it."""
    rows = build_rows(payload())
    assert len(rows["Replies"]) == 2           # header plus the one reply
    reply = rows["Replies"][1]
    assert reply[0] == "C001" and reply[1] == 1
    assert reply[2] == "ans.ahmad"
    assert reply[4] == "Agreed."


def test_the_reply_count_matches_the_replies_sheet():
    rows = build_rows(payload())
    counts = {r[0]: r[COMMENT_COLUMNS.index("Replies")] for r in rows["Comments"][1:]}
    actual = {}
    for r in rows["Replies"][1:]:
        actual[r[0]] = actual.get(r[0], 0) + 1
    for short_id, count in counts.items():
        assert count == actual.get(short_id, 0), short_id


def test_open_and_resolved_are_words_not_booleans():
    """A spreadsheet gets filtered by a person reading it."""
    rows = build_rows(payload())
    status = {r[0]: r[COMMENT_COLUMNS.index("Status")] for r in rows["Comments"][1:]}
    assert status == {"C001": "Open", "C002": "Resolved"}


def test_a_figure_reads_as_a_figure():
    rows = build_rows(payload())
    assert rows["Comments"][1][COMMENT_COLUMNS.index("Figure or table")] == "Figure 3"


def test_an_unnumbered_figure_says_what_it_can_rather_than_nothing():
    """When the number cannot be trusted the caption still identifies it, which
    is the fallback the multi-file work relies on."""
    assert _float_label({"kind": "figure", "number": None,
                         "caption": "A plot", "label": "fig:x"}) == 'Figure “A plot” (fig:x)'
    assert _float_label({"kind": "table", "number": None,
                         "caption": None, "label": None}) == "Table (unnumbered)"
    assert _float_label(None) == ""


def test_an_author_with_only_an_email_is_still_named():
    rows = build_rows(payload())
    assert rows["Comments"][2][COMMENT_COLUMNS.index("Author")] == "someone@example.com"


def test_tracked_changes_get_their_own_sheet():
    rows = build_rows(payload())
    assert len(rows["Tracked changes"]) == 2
    assert rows["Tracked changes"][1][0] == "T001"
    assert rows["Tracked changes"][1][4] == "Insertion"


def test_an_empty_project_still_produces_usable_sheets():
    rows = build_rows({"threads": {}, "comments": [], "tracked_changes": []})
    for name, table in rows.items():
        assert len(table) == 1, f"{name} should be headers only"
        assert table[0], f"{name} has no headers"


def test_every_row_matches_its_header_width():
    """A short row silently shifts every column after it."""
    rows = build_rows(payload())
    for name, table in rows.items():
        width = len(table[0])
        for i, row in enumerate(table[1:], start=2):
            assert len(row) == width, f"{name} row {i} has {len(row)} of {width}"


# --- characters a comment can carry that a spreadsheet cannot ---------------

HOSTILE = "bell\x07 vtab\x0b ff\x0c esc\x1b noncharacter￾ and ]]> <&>"


def _well_formed(path):
    import xml.parsers.expat
    import zipfile
    bad = []
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if name.endswith((".xml", ".rels")):
                parser = xml.parsers.expat.ParserCreate()
                try:
                    parser.Parse(z.read(name), True)
                except xml.parsers.expat.ExpatError as e:
                    bad.append(f"{name}: {e}")
    return bad


def test_a_comment_with_characters_xml_forbids_still_makes_a_spreadsheet(tmp_path):
    """openpyxl raised IllegalCharacterError on a control character, and that
    ended the whole export. Text pasted from Word can carry a vertical tab."""
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook
    from overleaf_comments_export.sheets import write_xlsx

    data = payload()
    data["threads"]["t1"]["messages"][0]["content"] = HOSTILE
    out = tmp_path / "comments.xlsx"
    write_xlsx(data, out)

    assert _well_formed(out) == []
    texts = [c.value for row in load_workbook(out)["Comments"].iter_rows() for c in row if isinstance(c.value, str)]
    kept = next(t for t in texts if "bell" in t)
    assert kept == "bell vtab ff esc noncharacter and ]]> <&>"


def test_a_cell_is_cut_to_what_excel_can_hold(tmp_path):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook
    from overleaf_comments_export.sheets import write_xlsx

    data = payload()
    data["threads"]["t1"]["messages"][0]["content"] = "x" * 40000
    out = tmp_path / "comments.xlsx"
    write_xlsx(data, out)
    longest = max(len(c.value) for row in load_workbook(out)["Comments"].iter_rows()
                  for c in row if isinstance(c.value, str))
    assert longest == 32767


def test_the_export_survives_such_a_comment_with_the_spreadsheet_on(tmp_path, monkeypatch):
    """The real failure: no files at all, because one cell raised."""
    pytest.importorskip("openpyxl")
    from overleaf_comments_export import export as export_mod
    from tests.test_export_wiring import FakeClient

    class Hostile(FakeClient):
        def get_threads(self, project_id):
            threads = super().get_threads(project_id)
            threads["t1"]["messages"][0]["content"] = HOSTILE
            return threads

    monkeypatch.setattr(export_mod, "OverleafClient", Hostile)
    export_mod.run_export(project_url="https://www.overleaf.com/project/" + "a" * 24,
                          out_dir=tmp_path, write_xlsx_sheet=True)
    sheets = list(tmp_path.rglob("comments.xlsx"))
    assert sheets, sorted(p.name for p in tmp_path.rglob("*"))
    assert _well_formed(sheets[0]) == []
